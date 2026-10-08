#!/usr/bin/env python3
r"""docx 读写工具：盘点结构、填充段落/表格、插图、另存。

依赖：python-docx

用法：
    python docx_io.py inspect <file.docx>
    python docx_io.py copy <模板.docx> <输出.docx>
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches


# ---------- 读取 ----------

def load(path):
    return Document(str(path))


def all_paragraphs(doc):
    """返回全部段落对象列表（含嵌套表格内）。"""
    return [obj for kind, obj in iter_body_items(doc) if kind == "p"]


def para_text(p) -> str:
    return "".join(r.text for r in p.runs).strip()


def iter_paragraphs(doc):
    """产出 (index, paragraph)，含表格/嵌套表格内的段落。"""
    i = 0
    for kind, obj in iter_body_items(doc):
        if kind == "p":
            yield i, obj
            i += 1


def iter_body_items(doc):
    """按文档顺序递归产出 ('p', Paragraph) / ('tbl', Table)。

    真实报告常整篇套在嵌套表格里，doc.paragraphs/doc.tables 会漏掉，
    必须递归遍历 body。
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = doc.element.body
    for e in body.iter():
        if e.tag == qn("w:p"):
            yield ("p", Paragraph(e, doc))
        elif e.tag == qn("w:tbl"):
            yield ("tbl", Table(e, doc))


def find_paragraphs(doc, keyword: str):
    """返回文本包含 keyword 的 (index, paragraph) 列表。"""
    return [(i, p) for i, p in iter_paragraphs(doc) if keyword in para_text(p)]


def has_formula(paragraph) -> bool:
    return paragraph._p.findall(
        ".//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath") != []


def count_formulas(doc) -> int:
    return len(doc.element.body.findall(
        ".//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"))


def table_shape(table):
    return len(table.rows), len(table.columns)


# ---------- 写入 ----------

def set_paragraph_text(paragraph, text: str, keep_format: bool = True):
    """替换段落文本，尽量保留首个 run 的格式。"""
    if keep_format and paragraph.runs:
        first = paragraph.runs[0]
        for r in paragraph.runs[1:]:
            r._element.getparent().remove(r._element)
        first.text = text
    else:
        for r in list(paragraph.runs):
            r._element.getparent().remove(r._element)
        paragraph.add_run(text)


def append_run(paragraph, text: str):
    return paragraph.add_run(text)


def insert_paragraph_after(paragraph, text: str = "", style=None):
    """在指定段落后插入新段落，返回新段落。"""
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    np = Paragraph(new_p, paragraph._parent)
    if text:
        np.add_run(text)
    if style is not None:
        np.style = style
    return np


def fill_cell(table, row: int, col: int, text: str, clear: bool = True):
    cell = table.cell(row, col)
    if clear:
        for p in cell.paragraphs:
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
        cell.paragraphs[0].add_run(text)
    else:
        cell.paragraphs[-1].add_run(text)
    return cell


def fill_table_grid(table, rows, start_row: int = 0, start_col: int = 0):
    """用二维列表 rows 填充表格（None 跳过）。"""
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            if val is None:
                continue
            fill_cell(table, start_row + i, start_col + j, str(val))


def add_picture_to_paragraph(paragraph, image_path, width_inches=None):
    run = paragraph.add_run()
    if width_inches:
        run.add_picture(str(image_path), width=Inches(width_inches))
    else:
        run.add_picture(str(image_path))
    return run


def add_centered_picture(doc, image_path, width_inches=None, caption=None):
    """插入居中图片；图题在上方居中。返回 (caption_para|None, pic_para)。"""
    cap_para = None
    if caption:
        cap_para = doc.add_paragraph(caption)
        cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_para = doc.add_paragraph()
    pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_picture_to_paragraph(pic_para, image_path, width_inches)
    return cap_para, pic_para


# ---------- 公式（OMML）与文字混排 ----------

MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def find_oMaths(paragraph):
    """返回段落内的 oMath 元素（文档顺序）。"""
    return list(paragraph._p.iter("{%s}oMath" % MATH_NS))


def fill_omml_placeholder(paragraph, idx: int, latex: str, size_pt=None):
    """把段落里第 idx 个已有公式（骨架占位）换成 latex 生成的公式，

    位置与周边文字保持不变。python 里 idx 从 0 开始。
    """
    import omml
    from docx.oxml import parse_xml

    ms = find_oMaths(paragraph)
    if idx >= len(ms):
        raise IndexError(f"段落里只有 {len(ms)} 个公式，取不到第 {idx} 个")
    old = ms[idx]
    new = parse_xml(omml.inline(latex, size_pt))
    old.getparent().replace(old, new)
    return new


def style_run(run, size_pt=10.5, bold=False, cjk="宋体", latin="Times New Roman"):
    """按课程规范设置 run 字体（中文宋体、西文 Times New Roman）。"""
    from docx.oxml import OxmlElement
    from docx.shared import Pt

    run.font.size = Pt(size_pt)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)
    return run


def write_mixed(paragraph, parts, template_run=None, size_pt=10.5,
                cjk="宋体", latin="Times New Roman", clear=True, clear_math=True):
    """在同一个段落里混排「文字 + 行内公式」。

    parts: list[str | ("math", latex) | ("block", latex)]
           str           -> 普通文字
           ("math", tex) -> 行内公式（不要用 Unicode 字符，写 LaTeX）
           ("block",tex) -> 独立居中公式
    template_run: 提供格式模板的 run（复制其 rPr，保留骨架字体设置）
    clear / clear_math: 写入前是否清掉原有文字 / 原有公式（默认都清，避免与
    骨架里的遗留占位公式重复）
    返回 paragraph。
    """
    import omml
    from copy import deepcopy
    from docx.oxml import parse_xml

    if clear:
        for r in list(paragraph.runs):
            r._element.getparent().remove(r._element)
    if clear_math:
        for om in find_oMaths(paragraph):
            om.getparent().remove(om)
    for part in parts:
        if isinstance(part, str):
            run = paragraph.add_run(part)
            rpr = None if template_run is None else template_run._element.find(qn("w:rPr"))
            if rpr is not None:
                old = run._element.find(qn("w:rPr"))
                if old is not None:
                    run._element.remove(old)
                run._element.insert(0, deepcopy(rpr))
            else:
                style_run(run, size_pt, cjk=cjk, latin=latin)
        else:
            kind, latex = part[0], part[1]
            frag = omml.block(latex, size_pt) if kind == "block" else omml.inline(latex, size_pt)
            paragraph._p.append(parse_xml(frag))
    return paragraph


def set_cell_math(cell, latex: str, size_pt=None, center: bool = True):
    """把单元格内容设成一条公式（表头符号用）。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import parse_xml
    import omml

    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for om in find_oMaths(p):        # 清掉骨架遗留的占位公式，避免重复
        om.getparent().remove(om)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p._p.append(parse_xml(omml.inline(latex, size_pt)))
    return p


def append_math(paragraph, latex: str, size_pt=None):
    """在段落末尾追加一条行内公式。"""
    from docx.oxml import parse_xml
    import omml

    paragraph._p.append(parse_xml(omml.inline(latex, size_pt)))
    return paragraph


def copy_template(template: str, out: str) -> str:
    import shutil
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template, out)
    return out


def save(doc, path):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


# ---------- CLI ----------

def _main(argv):
    if not argv or argv[0] in ("-h", "--help"):
        print(__doc__)
        return 0
    if argv[0] == "inspect":
        doc = load(argv[1])
        ps = all_paragraphs(doc)
        print(f"段落数: {len(ps)}  表格数: {len(doc.element.body.findall('.//' + qn('w:tbl')))}  公式数: {count_formulas(doc)}")
        for kind, obj in iter_body_items(doc):
            if kind == "p":
                t = para_text(obj)
                if t:
                    print(f"[p] {t[:80]}")
            else:
                print(f"[tbl] {table_shape(obj)}")
        return 0
    if argv[0] == "copy":
        print(copy_template(argv[1], argv[2]))
        return 0
    print(f"未知参数: {argv[0]}", file=sys.stderr)
    return 2


if __name__ == "__main__":
    raise SystemExit(_main(sys.argv[1:]))
